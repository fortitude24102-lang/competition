from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "word"
SPEC = ROOT / "docs" / "superpowers" / "specs" / "2026-09-02-dual-issue-ooo-contract-design.md"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5D6B7A"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
GRID = "B8C4D1"
WHITE = "FFFFFF"
CAUTION = "7A5A00"

PLANS = [
    {
        "source": ROOT / "docs" / "superpowers" / "plans" / "2026-09-02-ooo-backend-lead-plan.md",
        "output": OUT / "01_负责人_双发射乱序后端与集成计划.docx",
        "role": "负责人 / 成员1",
        "branch": "codex/ooo-backend-lead",
        "subtitle": "双译码、重命名、ROB、双发射、LSQ、精确异常与最终集成",
        "kicker": "高性能 RISC-V CPU · 工程执行手册",
        "callout": "你是唯一集成人。先冻结共享接口，再让两名组员从同一契约提交建立分支；只有独立验收通过的模块才能合入。",
        "append_contract": True,
    },
    {
        "source": ROOT / "docs" / "superpowers" / "plans" / "2026-09-02-ooo-frontend-member2-plan.md",
        "output": OUT / "02_成员2_双取指分支预测与ICache新手计划.docx",
        "role": "成员2 / 前端负责人",
        "branch": "codex/ooo-frontend-member2",
        "subtitle": "BTB、GShare、RAS、16KB I-Cache 与双取指 Frontend",
        "kicker": "零基础可执行版 · 每一步都有输入、输出和验收",
        "callout": "只改 frontend 目录。看不懂后端没关系：严格遵守 valid/ready/fire 和固定端口，专项测试通过后停止修改并把提交哈希交给负责人。",
        "append_contract": False,
    },
    {
        "source": ROOT / "docs" / "superpowers" / "plans" / "2026-09-02-ooo-exu-cache-member3-plan.md",
        "output": OUT / "03_成员3_RV32M与DCache新手计划.docx",
        "role": "成员3 / 执行与缓存负责人",
        "branch": "codex/ooo-memory-member3",
        "subtitle": "RV32M MulDiv、16KB 两路写回式 D-Cache 与独立验收",
        "kicker": "零基础可执行版 · 固定运算表、状态机和测试顺序",
        "callout": "只改 exu 和 cache 目录。tag 必须原样返回，MMIO、LSQ、符号扩展和 SoC 集成都不属于你的任务。专项验收后停止修改。",
        "append_contract": False,
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), GRID)


def set_table_geometry(table, widths_dxa: list[int]) -> None:
    total = sum(widths_dxa)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    set_table_borders(table)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def set_run_font(run, size=None, bold=None, italic=None, color=None, mono=False) -> None:
    name = "Consolas" if mono else "Calibri"
    east = "Microsoft YaHei"
    run.font.name = name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_inline(paragraph, text: str, size=None, color=None) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=9.5 if size is None else size, color=DARK_BLUE, mono=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, color=color)


def shade_paragraph(paragraph, fill: str, border: str | None = None) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    # Keep each field in its own run. LibreOffice can update these fields
    # reliably when the field instruction is surrounded by begin/separate/end.
    for field, suffix in ((" PAGE ", " 页 / 共 "), (" NUMPAGES ", " 页")):
        field_run = paragraph.add_run()
        set_run_font(field_run, size=9, color=MUTED)
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = field
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        field_run._r.extend([begin, instr, separate, end])
        suffix_run = paragraph.add_run(suffix)
        set_run_font(suffix_run, size=9, color=MUTED)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor.from_string(NAVY)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.1)
    code.paragraph_format.space_before = Pt(3)
    code.paragraph_format.space_after = Pt(5)
    code.paragraph_format.line_spacing = 1.0


def configure_section(doc: Document, running_label: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(running_label)
    set_run_font(run, size=9, bold=True, color=MUTED)

    footer = section.footer
    add_page_number(footer.paragraphs[0])


def add_cover(doc: Document, meta: dict) -> None:
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_after = Pt(8)
    run = kicker.add_run(meta["kicker"])
    set_run_font(run, size=11, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run(meta["source"].read_text(encoding="utf-8").splitlines()[0].lstrip("# "))
    set_run_font(run, size=28, bold=True, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run(meta["subtitle"])
    set_run_font(run, size=13.5, color=MUTED)

    table = doc.add_table(rows=3, cols=2)
    values = (("角色", meta["role"]), ("工作分支", meta["branch"]), ("版本", "V1.0 · 2026-09-02"))
    for row, (label, value) in zip(table.rows, values):
        row.cells[0].text = label
        row.cells[1].text = value
        set_cell_shading(row.cells[0], LIGHT_BLUE)
        for idx, cell in enumerate(row.cells):
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_run_font(run, size=10.5, bold=(idx == 0), color=NAVY)
    repeat_table_header(table.rows[0])
    set_table_geometry(table, [1700, 7660])

    doc.add_paragraph()
    callout = doc.add_paragraph()
    callout.paragraph_format.left_indent = Inches(0.18)
    callout.paragraph_format.right_indent = Inches(0.12)
    callout.paragraph_format.space_before = Pt(8)
    callout.paragraph_format.space_after = Pt(8)
    add_inline(callout, "执行原则：" + meta["callout"], size=11, color=NAVY)
    shade_paragraph(callout, LIGHT_GRAY, BLUE)

    footer_note = doc.add_paragraph()
    footer_note.paragraph_format.space_before = Pt(34)
    footer_note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = footer_note.add_run("项目：盘古 MES2L676-100HP · Chisel RV32IM 双发射乱序 CPU")
    set_run_font(run, size=9.5, color=MUTED)
    footer_note.add_run().add_break(WD_BREAK.PAGE)


def table_widths(column_count: int) -> list[int]:
    if column_count == 2:
        return [2300, 7060]
    if column_count == 3:
        return [1800, 3780, 3780]
    if column_count == 4:
        return [1400, 2400, 2780, 2780]
    base = 9360 // column_count
    widths = [base] * column_count
    widths[-1] += 9360 - sum(widths)
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]]) -> None:
    if len(rows) < 2:
        return
    header = rows[0]
    body = rows[2:] if all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in rows[1]) else rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    for idx, text in enumerate(header):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        add_inline(p, text.strip(), size=9.2, color=NAVY)
        for run in p.runs:
            run.bold = True
    repeat_table_header(table.rows[0])

    for row_data in body:
        row = table.add_row()
        for idx in range(len(header)):
            text = row_data[idx].strip() if idx < len(row_data) else ""
            p = row.cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_inline(p, text, size=9.2, color=None)
    set_table_geometry(table, table_widths(len(header)))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def parse_table_line(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def render_markdown(doc: Document, text: str, skip_first_h1=True) -> None:
    lines = text.splitlines()
    i = 0
    in_code = False
    code_lines: list[str] = []
    first_h1_skipped = False
    task_seen = False

    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph(style="Code Block")
                for idx, code_line in enumerate(code_lines):
                    if idx:
                        p.add_run().add_break()
                    run = p.add_run(code_line)
                    set_run_font(run, size=8.5, color=NAVY, mono=True)
                shade_paragraph(p, "F3F6F9")
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(parse_table_line(lines[i]))
                i += 1
            add_markdown_table(doc, rows)
            continue

        if not line.strip() or line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            if skip_first_h1 and not first_h1_skipped:
                first_h1_skipped = True
            else:
                doc.add_paragraph(line[2:].strip(), style="Heading 1")
            i += 1
            continue
        if line.startswith("## "):
            doc.add_paragraph(line[3:].strip(), style="Heading 1")
            i += 1
            continue
        if line.startswith("### "):
            heading = line[4:].strip()
            p = doc.add_paragraph(style="Heading 2")
            if heading.startswith("Task ") and task_seen:
                p.paragraph_format.page_break_before = True
            if heading.startswith("Task "):
                task_seen = True
            add_inline(p, heading)
            i += 1
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            add_inline(p, line[2:].strip(), size=9.5, color=MUTED)
            shade_paragraph(p, LIGHT_GRAY, BLUE)
            i += 1
            continue

        match = re.match(r"^- \[ \] (.+)$", line)
        if match:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, "☐ " + match.group(1))
            i += 1
            continue
        match = re.match(r"^- (.+)$", line)
        if match:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, match.group(1))
            i += 1
            continue
        match = re.match(r"^\d+\. (.+)$", line)
        if match:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, match.group(1))
            i += 1
            continue

        paragraph_lines = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if not nxt.strip() or nxt.startswith(("#", "- ", "> ", "```", "|")) or re.match(r"^\d+\. ", nxt):
                break
            paragraph_lines.append(nxt.strip())
            i += 1
        p = doc.add_paragraph()
        add_inline(p, " ".join(paragraph_lines))


def build_doc(meta: dict) -> None:
    doc = Document()
    configure_styles(doc)
    configure_section(doc, f"双发射乱序 CPU 实施计划 · {meta['role']}")
    add_cover(doc, meta)

    source_text = meta["source"].read_text(encoding="utf-8")
    render_markdown(doc, source_text, skip_first_h1=True)

    if meta["append_contract"]:
        p = doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        contract_title = doc.add_paragraph("附录：三人共用接口与分工契约", style="Heading 1")
        contract_title.paragraph_format.page_break_before = False
        render_markdown(doc, SPEC.read_text(encoding="utf-8"), skip_first_h1=True)

    doc.core_properties.title = meta["source"].read_text(encoding="utf-8").splitlines()[0].lstrip("# ")
    doc.core_properties.subject = "盘古 MES2L676-100HP 双发射乱序 CPU 三人实施计划"
    doc.core_properties.author = "项目组"
    doc.core_properties.keywords = "RISC-V, Chisel, 双发射, 乱序执行, CoreMark"
    doc.core_properties.comments = "由项目 Markdown 计划生成的 Word 执行手册"

    meta["output"].parent.mkdir(parents=True, exist_ok=True)
    doc.save(meta["output"])


def main() -> None:
    for meta in PLANS:
        build_doc(meta)
        print(meta["output"])


if __name__ == "__main__":
    main()
