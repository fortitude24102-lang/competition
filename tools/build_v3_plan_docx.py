"""Build the V3 Word plan from its reviewed Markdown and interface appendix."""
from pathlib import Path
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
PLAN = ROOT / "docs/efinix_2d_gpu/v3_two_week_plan_20261001.md"
SPEC = ROOT / "docs/efinix_2d_gpu/v3_interface_design_20261001.md"
OUTPUT = ROOT / "docs/word/Efinix_2D图像渲染V3两周开发计划书_网口交互与GPU优化版.docx"


def inline(paragraph, text):
    # Keep link destinations visible in Word, so the document works offline.
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if not part:
            continue
        run = paragraph.add_run(part.strip("*`") if part.startswith(("**", "`")) else part)
        run.bold = part.startswith("**")


def add_markdown(doc, path, appendix=False):
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        if line.startswith("# "):
            doc.add_paragraph("附录 V3 架构与固定接口" if appendix else line[2:],
                              "Heading 1" if appendix else "Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], "Heading 1")
        elif line.startswith("### "):
            doc.add_paragraph(line[4:], "Heading 2")
        else:
            bullet = line.startswith("- ")
            paragraph = doc.add_paragraph(style="List Bullet" if bullet else "Normal")
            inline(paragraph, line[2:] if bullet else line)


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for name, size in (("Normal", 11), ("List Bullet", 11), ("Title", 22),
                       ("Heading 1", 15), ("Heading 2", 12)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.line_spacing = 1.12
        style.paragraph_format.space_after = Pt(6)
        if name.startswith("Heading"):
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.keep_with_next = True
        # Built-in Title must not inherit a border or theme text color.
        ppr = style.element.find(qn("w:pPr"))
        if ppr is not None:
            for border in list(ppr.findall(qn("w:pBdr"))):
                ppr.remove(border)
        rpr = style.element.find(qn("w:rPr"))
        if rpr is not None:
            color = rpr.find(qn("w:color"))
            if color is not None:
                for attribute in ("themeColor", "themeTint", "themeShade"):
                    color.attrib.pop(qn("w:" + attribute), None)
    add_markdown(doc, PLAN)
    doc.add_page_break()
    add_markdown(doc, SPEC, appendix=True)
    # Page numbers help a team refer to the same contract during integration.
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    doc.core_properties.title = "AetherGX V3 两周开发计划书"
    doc.core_properties.subject = "GPU优化与网口互动的两条线交付计划"
    doc.core_properties.author = "AetherGX 项目组"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
